from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import os
import shutil
import subprocess
import tarfile
import tempfile
import time
import zipfile
from collections import defaultdict
from pathlib import Path, PurePosixPath
from typing import Any, Iterable

from .operator import BLOCKED_NAMES, BLOCKED_PARTS, MINIMAL_COMMAND_ENV, PathPolicy, _trim, blocked_name


MAX_DOCUMENT_BYTES = 250_000_000
MAX_DOCUMENT_PAGES = 100
MAX_OCR_PAGES = 30
MAX_WORKBOOK_SHEETS = 100
MAX_WORKBOOK_ROWS = 10_000
MAX_WORKBOOK_CELLS = 100_000
MAX_PRESENTATION_SLIDES = 500
MAX_ARCHIVE_FILES = 2_000
MAX_ARCHIVE_MEMBER_BYTES = 1_000_000_000
MAX_ARCHIVE_TOTAL_BYTES = 5_000_000_000
MAX_SEARCH_RESULTS = 200
MAX_DUPLICATE_FILES = 50_000
MAX_TRANSCRIPTION_SECONDS = 7_200
OUTPUT_ROOT = Path("/home/darren/Documents/Little-Oogway/workbench-output")
WHISPER_CACHE = Path("/mnt/storage/models/faster-whisper")
DOCUMENT_EXTENSIONS = {
    ".pdf", ".docx", ".odt", ".epub", ".html", ".htm", ".md", ".markdown", ".txt",
    ".xlsx", ".ods", ".csv", ".tsv", ".pptx", ".odp", ".doc", ".xls", ".ppt", ".rtf",
}
LEGACY_OFFICE_TARGETS = {".doc": ".docx", ".rtf": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
ARCHIVE_EXTENSIONS = {".zip", ".7z", ".tar", ".tgz", ".gz", ".bz2", ".xz"}
RIPGREP_CANDIDATES = (Path("/home/darren/.local/bin/rg"), Path("/usr/bin/rg"))


def _safe_output_name(source: Path, suffix: str) -> Path:
    fingerprint = hashlib.sha256(str(source).encode("utf-8")).hexdigest()[:12]
    stem = "".join(ch if ch.isalnum() or ch in "-_" else "-" for ch in source.stem)[:80]
    return OUTPUT_ROOT / f"{stem}-{fingerprint}{suffix}"


def _bounded_file(policy: PathPolicy, value: str, *, max_bytes: int = MAX_DOCUMENT_BYTES) -> Path:
    target = policy.resolve(value)
    if not target.is_file():
        raise ValueError("path is not a regular file")
    if target.stat().st_size > max_bytes:
        raise ValueError(f"file exceeds the {max_bytes:,}-byte limit")
    return target


def _clean_text(value: str) -> str:
    lines = [line.rstrip() for line in value.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    output: list[str] = []
    blank = False
    for line in lines:
        if line.strip():
            output.append(line)
            blank = False
        elif not blank:
            output.append("")
            blank = True
    return "\n".join(output).strip()


def _sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _ripgrep_path() -> Path:
    for candidate in RIPGREP_CANDIDATES:
        if candidate.is_file():
            return candidate
    raise RuntimeError("ripgrep is not installed in an approved executable location")


class DocumentRuntime:
    def __init__(self, policy: PathPolicy | None = None) -> None:
        self.policy = policy or PathPolicy()
        self._ocr_engine: Any | None = None

    def _ocr_image(self, source: Path) -> tuple[str, list[float]]:
        if self._ocr_engine is None:
            from rapidocr_onnxruntime import RapidOCR

            self._ocr_engine = RapidOCR()
        results, _elapsed = self._ocr_engine(source)
        if not results:
            return "", []
        text = []
        scores = []
        for item in results:
            if len(item) >= 3:
                text.append(str(item[1]))
                scores.append(float(item[2]))
        return "\n".join(text), scores

    @staticmethod
    def _tabular_row(values: Iterable[Any]) -> str:
        return " | ".join(
            str(value).replace("\n", " ").replace("|", "\\|") if value is not None else ""
            for value in values
        )

    def _extract_xlsx(self, source: Path) -> tuple[str, dict[str, Any]]:
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        blocks: list[str] = []
        cells = 0
        rows = 0
        sheets = workbook.worksheets[:MAX_WORKBOOK_SHEETS]
        try:
            for sheet in sheets:
                blocks.append(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    if rows >= MAX_WORKBOOK_ROWS or cells >= MAX_WORKBOOK_CELLS:
                        break
                    values = list(row)[: MAX_WORKBOOK_CELLS - cells]
                    blocks.append(self._tabular_row(values))
                    rows += 1
                    cells += len(values)
        finally:
            workbook.close()
        return "\n".join(blocks), {
            "format": "xlsx", "sheets": len(sheets), "rows_processed": rows,
            "cells_processed": cells, "bounded": rows >= MAX_WORKBOOK_ROWS or cells >= MAX_WORKBOOK_CELLS,
        }

    def _extract_delimited(self, source: Path, delimiter: str | None = None) -> tuple[str, dict[str, Any]]:
        rows = 0
        cells = 0
        blocks: list[str] = []
        with source.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
            if delimiter is None:
                sample = stream.read(32_768)
                stream.seek(0)
                try:
                    delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                except csv.Error:
                    delimiter = ","
            reader = csv.reader(stream, delimiter=delimiter)
            for row in reader:
                if rows >= MAX_WORKBOOK_ROWS or cells >= MAX_WORKBOOK_CELLS:
                    break
                values = row[: MAX_WORKBOOK_CELLS - cells]
                blocks.append(self._tabular_row(values))
                rows += 1
                cells += len(values)
        return "\n".join(blocks), {
            "format": "tsv" if delimiter == "\t" else "csv", "delimiter": delimiter,
            "rows_processed": rows, "cells_processed": cells,
            "bounded": rows >= MAX_WORKBOOK_ROWS or cells >= MAX_WORKBOOK_CELLS,
        }

    def _extract_ods(self, source: Path) -> tuple[str, dict[str, Any]]:
        from odf import teletype
        from odf.opendocument import load
        from odf.table import Table, TableCell, TableRow

        document = load(str(source))
        blocks: list[str] = []
        cells = 0
        rows = 0
        sheets = document.getElementsByType(Table)[:MAX_WORKBOOK_SHEETS]
        for sheet in sheets:
            title = sheet.getAttribute("name") or f"Sheet {len(blocks) + 1}"
            blocks.append(f"## Sheet: {title}")
            for row in sheet.getElementsByType(TableRow):
                repeat_rows = min(int(row.getAttribute("numberrowsrepeated") or 1), MAX_WORKBOOK_ROWS - rows)
                values: list[str] = []
                for cell in row.getElementsByType(TableCell):
                    value = teletype.extractText(cell)
                    if not value:
                        value = cell.getAttribute("value") or cell.getAttribute("stringvalue") or ""
                    repeat_cells = min(
                        int(cell.getAttribute("numbercolumnsrepeated") or 1),
                        MAX_WORKBOOK_CELLS - cells - len(values),
                    )
                    values.extend([value] * max(0, repeat_cells))
                    if cells + len(values) >= MAX_WORKBOOK_CELLS:
                        break
                for _ in range(max(0, repeat_rows)):
                    blocks.append(self._tabular_row(values))
                    rows += 1
                    cells += len(values)
                    if rows >= MAX_WORKBOOK_ROWS or cells >= MAX_WORKBOOK_CELLS:
                        break
                if rows >= MAX_WORKBOOK_ROWS or cells >= MAX_WORKBOOK_CELLS:
                    break
        return "\n".join(blocks), {
            "format": "ods", "sheets": len(sheets), "rows_processed": rows,
            "cells_processed": cells, "bounded": rows >= MAX_WORKBOOK_ROWS or cells >= MAX_WORKBOOK_CELLS,
        }

    def _extract_pptx(self, source: Path) -> tuple[str, dict[str, Any]]:
        from pptx import Presentation

        presentation = Presentation(source)
        slides = list(presentation.slides)[:MAX_PRESENTATION_SLIDES]
        blocks: list[str] = []
        for index, slide in enumerate(slides, start=1):
            parts = [f"## Slide {index}"]
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        parts.append(text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        parts.append(self._tabular_row(cell.text for cell in row.cells))
            blocks.append("\n\n".join(parts))
        return "\n\n".join(blocks), {
            "format": "pptx", "slides": len(presentation.slides), "slides_processed": len(slides),
            "bounded": len(presentation.slides) > len(slides),
        }

    def _extract_odp(self, source: Path) -> tuple[str, dict[str, Any]]:
        from odf import teletype
        from odf.draw import Page
        from odf.opendocument import load
        from odf.text import H, P

        document = load(str(source))
        pages = document.getElementsByType(Page)
        selected = pages[:MAX_PRESENTATION_SLIDES]
        blocks: list[str] = []
        for index, page in enumerate(selected, start=1):
            text_nodes = list(page.getElementsByType(H)) + list(page.getElementsByType(P))
            content = [teletype.extractText(node).strip() for node in text_nodes]
            blocks.append("\n\n".join([f"## Slide {index}"] + [item for item in content if item]))
        return "\n\n".join(blocks), {
            "format": "odp", "slides": len(pages), "slides_processed": len(selected),
            "bounded": len(pages) > len(selected),
        }

    def _convert_legacy_office(self, source: Path) -> tuple[Path, tempfile.TemporaryDirectory[str]]:
        target_suffix = LEGACY_OFFICE_TARGETS[source.suffix.casefold()]
        OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
        temporary = tempfile.TemporaryDirectory(prefix=".office-convert-", dir=OUTPUT_ROOT)
        destination = Path(temporary.name) / "output"
        destination.mkdir()
        profile = Path(temporary.name) / "profile"
        staged_source = Path(temporary.name) / f"source{source.suffix.casefold()}"
        shutil.copy2(source, staged_source)
        libreoffice = next(
            (path for path in (Path("/snap/bin/libreoffice"), Path("/usr/bin/libreoffice"), Path("/usr/bin/soffice")) if path.exists()),
            None,
        )
        if libreoffice is None:
            temporary.cleanup()
            raise RuntimeError("LibreOffice is not installed")
        env = dict(MINIMAL_COMMAND_ENV)
        env.update({"LANG": "C.UTF-8", "SAL_USE_VCLPLUGIN": "svp"})
        completed = subprocess.run(
            [str(libreoffice), f"-env:UserInstallation=file://{profile}", "--headless", "--nologo", "--nodefault", "--nolockcheck", "--norestore", "--convert-to", target_suffix.lstrip("."), "--outdir", str(destination), str(staged_source)],
            capture_output=True, text=True, timeout=180, env=env,
        )
        outputs = list(destination.glob(f"*{target_suffix}"))
        if completed.returncode or len(outputs) != 1:
            temporary.cleanup()
            message = (completed.stderr or completed.stdout).strip()[:1000]
            raise RuntimeError(message or "LibreOffice conversion failed")
        return outputs[0], temporary

    def ocr(self, path: str, max_pages: int = 10) -> dict[str, Any]:
        source = _bounded_file(self.policy, path)
        page_limit = min(max(int(max_pages), 1), MAX_OCR_PAGES)
        pages: list[Path]
        with tempfile.TemporaryDirectory(prefix="oogway-ocr-") as directory:
            temporary = Path(directory)
            if source.suffix.casefold() == ".pdf":
                prefix = temporary / "page"
                completed = subprocess.run(
                    [
                        "/usr/bin/pdftoppm", "-f", "1", "-l", str(page_limit),
                        "-r", "200", "-jpeg", str(source), str(prefix),
                    ],
                    capture_output=True,
                    text=True,
                    timeout=300,
                    env=MINIMAL_COMMAND_ENV,
                )
                if completed.returncode:
                    raise RuntimeError(completed.stderr.strip()[:1000] or "PDF rendering failed")
                pages = sorted(temporary.glob("page-*.jpg"))
            elif source.suffix.casefold() in IMAGE_EXTENSIONS:
                normalized = temporary / "page-1.jpg"
                completed = subprocess.run(
                    [
                        "/usr/bin/ffmpeg", "-hide_banner", "-loglevel", "error",
                        "-threads", "1", "-y", "-i", str(source), "-vf",
                        "scale=4096:4096:force_original_aspect_ratio=decrease",
                        "-frames:v", "1", "-pix_fmt", "yuvj420p", str(normalized),
                    ],
                    capture_output=True,
                    text=True,
                    timeout=90,
                    env=MINIMAL_COMMAND_ENV,
                )
                if completed.returncode or not normalized.is_file():
                    raise RuntimeError(completed.stderr.strip()[:1000] or "image normalization failed")
                pages = [normalized]
            else:
                raise ValueError("OCR accepts PDF or common image formats")
            if not pages:
                raise ValueError("OCR source produced no pages")
            page_results = []
            combined = []
            all_scores: list[float] = []
            for index, page in enumerate(pages, start=1):
                text, scores = self._ocr_image(page)
                combined.append(f"## Page {index}\n\n{text}".strip())
                all_scores.extend(scores)
                page_results.append({"page": index, "characters": len(text), "lines": len(text.splitlines())})
        content, truncated = _trim("\n\n".join(combined), 200_000)
        return {
            "source": str(source),
            "source_sha256": _sha256_path(source),
            "pages_processed": len(page_results),
            "pages": page_results,
            "mean_confidence": round(sum(all_scores) / len(all_scores), 4) if all_scores else None,
            "content": content,
            "truncated": truncated,
            "engine": "RapidOCR ONNX CPU",
        }

    def _extract(self, source: Path, *, ocr_if_needed: bool, max_pages: int) -> tuple[str, dict[str, Any]]:
        suffix = source.suffix.casefold()
        metadata: dict[str, Any] = {"format": suffix.lstrip(".") or "text"}
        if suffix in LEGACY_OFFICE_TARGETS:
            converted, temporary = self._convert_legacy_office(source)
            try:
                text, converted_metadata = self._extract(
                    converted, ocr_if_needed=ocr_if_needed, max_pages=max_pages
                )
                intermediate_format = converted.suffix.lstrip(".")
            finally:
                temporary.cleanup()
            converted_metadata.update(
                {
                    "format": suffix.lstrip("."),
                    "converted_via": "LibreOffice",
                    "intermediate_format": intermediate_format,
                }
            )
            return text, converted_metadata
        if suffix in {".txt", ".md", ".markdown"}:
            return source.read_text(encoding="utf-8", errors="replace"), metadata
        if suffix in {".html", ".htm"}:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(source.read_bytes(), "html.parser")
            for element in soup(["script", "style", "noscript", "template"]):
                element.decompose()
            metadata["title"] = soup.title.get_text(" ", strip=True) if soup.title else ""
            return soup.get_text("\n"), metadata
        if suffix == ".docx":
            from docx import Document

            document = Document(source)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            metadata["paragraphs"] = len(document.paragraphs)
            metadata["tables"] = len(document.tables)
            return "\n\n".join(paragraphs), metadata
        if suffix == ".odt":
            from odf import teletype
            from odf.opendocument import load
            from odf.text import H, P

            document = load(str(source))
            blocks = [teletype.extractText(node) for node in document.getElementsByType(H)]
            blocks.extend(teletype.extractText(node) for node in document.getElementsByType(P))
            return "\n\n".join(blocks), metadata
        if suffix == ".xlsx":
            return self._extract_xlsx(source)
        if suffix == ".ods":
            return self._extract_ods(source)
        if suffix in {".csv", ".tsv"}:
            return self._extract_delimited(source, "\t" if suffix == ".tsv" else None)
        if suffix == ".pptx":
            return self._extract_pptx(source)
        if suffix == ".odp":
            return self._extract_odp(source)
        if suffix == ".epub":
            from bs4 import BeautifulSoup
            from ebooklib import ITEM_DOCUMENT, epub

            book = epub.read_epub(str(source), options={"ignore_ncx": True})
            blocks = []
            for item in book.get_items_of_type(ITEM_DOCUMENT):
                soup = BeautifulSoup(item.get_content(), "html.parser")
                for element in soup(["script", "style"]):
                    element.decompose()
                blocks.append(soup.get_text("\n"))
            metadata["sections"] = len(blocks)
            return "\n\n".join(blocks), metadata
        if suffix == ".pdf":
            import pdfplumber

            pages = []
            with pdfplumber.open(source) as document:
                total_pages = len(document.pages)
                selected = document.pages[: min(max_pages, total_pages)]
                for index, page in enumerate(selected, start=1):
                    pages.append(f"## Page {index}\n\n{page.extract_text(layout=True) or ''}")
            text = "\n\n".join(pages)
            metadata.update({"pages": total_pages, "pages_processed": len(pages)})
            if ocr_if_needed and len("".join(text.split())) < max(80, len(pages) * 40):
                ocr = self.ocr(str(source), max_pages=min(max_pages, MAX_OCR_PAGES))
                metadata["ocr_fallback"] = True
                metadata["ocr_engine"] = ocr["engine"]
                return str(ocr["content"]), metadata
            metadata["ocr_fallback"] = False
            return text, metadata
        raise ValueError("unsupported document format")

    def read(self, path: str, *, ocr_if_needed: bool = True, max_pages: int = 30) -> dict[str, Any]:
        source = _bounded_file(self.policy, path)
        if source.suffix.casefold() not in DOCUMENT_EXTENSIONS:
            raise ValueError("unsupported document, spreadsheet, or presentation format")
        page_limit = min(max(int(max_pages), 1), MAX_DOCUMENT_PAGES)
        text, metadata = self._extract(source, ocr_if_needed=ocr_if_needed, max_pages=page_limit)
        content, truncated = _trim(_clean_text(text), 200_000)
        return {
            "source": str(source),
            "sha256": _sha256_path(source),
            "bytes": source.stat().st_size,
            "metadata": metadata,
            "content": content,
            "truncated": truncated,
        }

    def convert_to_markdown(self, path: str, *, ocr_if_needed: bool = True) -> dict[str, Any]:
        source = _bounded_file(self.policy, path)
        text, metadata = self._extract(
            source, ocr_if_needed=ocr_if_needed, max_pages=MAX_DOCUMENT_PAGES
        )
        content = _clean_text(text)
        if len(content) > 5_000_000:
            raise ValueError("converted document exceeds the five-million-character output limit")
        source_hash = _sha256_path(source)
        OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
        output = _safe_output_name(source, ".md")
        header = (
            f"# Converted document\n\n"
            f"- Source: `{source}`\n"
            f"- SHA-256: `{source_hash}`\n"
            f"- Converted: {time.strftime('%Y-%m-%dT%H:%M:%S%z')}\n\n"
        )
        output.write_text(header + content + "\n", encoding="utf-8")
        return {"source": str(source), "output": str(output), "sha256": source_hash, "metadata": metadata}


class MediaRuntime:
    def __init__(self, policy: PathPolicy | None = None) -> None:
        self.policy = policy or PathPolicy()

    def inspect(self, path: str) -> dict[str, Any]:
        source = _bounded_file(self.policy, path, max_bytes=10_000_000_000)
        completed = subprocess.run(
            [
                "/usr/bin/ffprobe", "-v", "error", "-show_format", "-show_streams",
                "-of", "json", str(source),
            ],
            capture_output=True,
            text=True,
            timeout=60,
            env=MINIMAL_COMMAND_ENV,
        )
        if completed.returncode:
            raise ValueError(completed.stderr.strip()[:1000] or "FFprobe could not inspect this file")
        probe = json.loads(completed.stdout or "{}")
        image_metadata: dict[str, Any] = {}
        if source.suffix.casefold() in IMAGE_EXTENSIONS:
            try:
                from PIL import ExifTags, Image

                with Image.open(source) as image:
                    image_metadata.update({"width": image.width, "height": image.height, "mode": image.mode, "format": image.format})
                    exif = image.getexif()
                    image_metadata["exif"] = {
                        str(ExifTags.TAGS.get(key, key)): str(value)[:1000]
                        for key, value in exif.items()
                    }
            except (OSError, ValueError):
                image_metadata["image_decode_error"] = True
        return {
            "path": str(source),
            "bytes": source.stat().st_size,
            "sha256": _sha256_path(source),
            "modified": source.stat().st_mtime,
            "ffprobe": probe,
            "image": image_metadata,
        }


class ArchiveRuntime:
    def __init__(self, policy: PathPolicy | None = None) -> None:
        self.policy = policy or PathPolicy()

    @staticmethod
    def _validate_members(members: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
        checked = []
        total = 0
        for member in members:
            name = str(member["name"]).replace("\\", "/")
            path = PurePosixPath(name)
            size = int(member.get("size", 0))
            if path.is_absolute() or ".." in path.parts or not name.strip():
                raise ValueError(f"unsafe archive path: {name!r}")
            if member.get("link"):
                raise ValueError(f"archive links are not allowed: {name!r}")
            if size > MAX_ARCHIVE_MEMBER_BYTES:
                raise ValueError(f"archive member exceeds one gigabyte: {name!r}")
            total += size
            if total > MAX_ARCHIVE_TOTAL_BYTES:
                raise ValueError("archive expands beyond the five-gigabyte limit")
            checked.append({"name": name, "size": size, "directory": bool(member.get("directory"))})
            if len(checked) > MAX_ARCHIVE_FILES:
                raise ValueError("archive contains more than 2,000 entries")
        return checked

    def _members(self, source: Path) -> tuple[str, list[dict[str, Any]]]:
        suffix = source.suffix.casefold()
        if suffix == ".zip":
            with zipfile.ZipFile(source) as archive:
                members = [
                    {"name": item.filename, "size": item.file_size, "directory": item.is_dir(), "link": (item.external_attr >> 16) & 0o170000 == 0o120000}
                    for item in archive.infolist()
                ]
            return "zip", self._validate_members(members)
        if suffix == ".7z":
            import py7zr

            with py7zr.SevenZipFile(source, mode="r") as archive:
                members = [
                    {"name": item.filename, "size": item.uncompressed or 0, "directory": item.is_directory, "link": bool(getattr(item, "is_symlink", False))}
                    for item in archive.list()
                ]
            return "7z", self._validate_members(members)
        if tarfile.is_tarfile(source):
            with tarfile.open(source, mode="r:*") as archive:
                members = [
                    {"name": item.name, "size": item.size, "directory": item.isdir(), "link": item.issym() or item.islnk()}
                    for item in archive.getmembers()
                ]
            return "tar", self._validate_members(members)
        raise ValueError("supported archives are ZIP, 7z, TAR, TGZ, TAR.GZ, TAR.BZ2, and TAR.XZ")

    def inspect(self, path: str) -> dict[str, Any]:
        source = _bounded_file(self.policy, path, max_bytes=10_000_000_000)
        kind, members = self._members(source)
        return {
            "path": str(source), "format": kind, "entries": len(members),
            "source_sha256": _sha256_path(source),
            "uncompressed_bytes": sum(item["size"] for item in members),
            "members": members[:200], "members_truncated": len(members) > 200,
        }

    def extract(self, path: str) -> dict[str, Any]:
        source = _bounded_file(self.policy, path, max_bytes=10_000_000_000)
        kind, members = self._members(source)
        OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
        output = _safe_output_name(source, "-extracted")
        if output.exists():
            return {"source": str(source), "source_sha256": _sha256_path(source), "output": str(output), "format": kind, "entries": len(members), "reused": True}
        with tempfile.TemporaryDirectory(prefix="oogway-extract-", dir=OUTPUT_ROOT) as directory:
            staging = Path(directory)
            if kind == "zip":
                with zipfile.ZipFile(source) as archive:
                    for item in archive.infolist():
                        destination = staging / PurePosixPath(item.filename)
                        if item.is_dir():
                            destination.mkdir(parents=True, exist_ok=True)
                            continue
                        destination.parent.mkdir(parents=True, exist_ok=True)
                        with archive.open(item) as incoming, destination.open("xb") as outgoing:
                            shutil.copyfileobj(incoming, outgoing, 1024 * 1024)
            elif kind == "tar":
                with tarfile.open(source, mode="r:*") as archive:
                    for item in archive.getmembers():
                        destination = staging / PurePosixPath(item.name)
                        if item.isdir():
                            destination.mkdir(parents=True, exist_ok=True)
                            continue
                        incoming = archive.extractfile(item)
                        if incoming is None:
                            continue
                        destination.parent.mkdir(parents=True, exist_ok=True)
                        with incoming, destination.open("xb") as outgoing:
                            shutil.copyfileobj(incoming, outgoing, 1024 * 1024)
            else:
                import py7zr

                with py7zr.SevenZipFile(source, mode="r") as archive:
                    archive.extractall(path=staging)
            os.replace(staging, output)
        return {"source": str(source), "source_sha256": _sha256_path(source), "output": str(output), "format": kind, "entries": len(members), "reused": False}


class AudioRuntime:
    def __init__(self, policy: PathPolicy | None = None) -> None:
        self.policy = policy or PathPolicy()
        self._model: Any | None = None

    def transcribe(self, path: str, language: str = "") -> dict[str, Any]:
        source = _bounded_file(self.policy, path, max_bytes=5_000_000_000)
        probe = subprocess.run(
            ["/usr/bin/ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "default=nw=1:nk=1", str(source)],
            capture_output=True, text=True, timeout=30, env=MINIMAL_COMMAND_ENV,
        )
        try:
            duration = float(probe.stdout.strip())
        except ValueError as exc:
            raise ValueError("audio duration could not be determined") from exc
        if duration > MAX_TRANSCRIPTION_SECONDS:
            raise ValueError("audio exceeds the two-hour transcription limit")
        if self._model is None:
            from faster_whisper import WhisperModel

            WHISPER_CACHE.mkdir(parents=True, exist_ok=True)
            self._model = WhisperModel("base", device="cpu", compute_type="int8", cpu_threads=6, download_root=str(WHISPER_CACHE))
        segments, info = self._model.transcribe(
            str(source), language=language.strip() or None, vad_filter=True, beam_size=5
        )
        rows = []
        for segment in segments:
            rows.append({"start": round(segment.start, 3), "end": round(segment.end, 3), "text": segment.text.strip()})
            if len(rows) >= 5_000:
                break
        transcript, truncated = _trim(" ".join(item["text"] for item in rows), 200_000)
        return {
            "source": str(source), "source_sha256": _sha256_path(source), "duration_seconds": duration,
            "language": info.language, "language_probability": info.language_probability,
            "model": "faster-whisper base (CPU int8)", "segments": rows,
            "transcript": transcript, "truncated": truncated,
        }


class FileRuntime:
    def __init__(self, policy: PathPolicy | None = None) -> None:
        self.policy = policy or PathPolicy()

    def capabilities(self) -> dict[str, Any]:
        commands = {
            "ffmpeg": Path("/usr/bin/ffmpeg"),
            "ffprobe": Path("/usr/bin/ffprobe"),
            "git": Path("/usr/bin/git"),
            "pdftoppm": Path("/usr/bin/pdftoppm"),
            "ripgrep": next((path for path in RIPGREP_CANDIDATES if path.is_file()), RIPGREP_CANDIDATES[0]),
            "libreoffice": next(
                (path for path in (Path("/snap/bin/libreoffice"), Path("/usr/bin/libreoffice"), Path("/usr/bin/soffice")) if path.exists()),
                Path("/snap/bin/libreoffice"),
            ),
        }
        packages = [
            "bs4", "docx", "ebooklib", "faster_whisper", "odf", "pdfplumber",
            "openpyxl", "PIL", "pptx", "py7zr", "rapidocr_onnxruntime",
        ]
        return {
            "commands": {name: str(path) if path.is_file() else None for name, path in commands.items()},
            "python_packages": {name: importlib.util.find_spec(name) is not None for name in packages},
            "document_formats": sorted(DOCUMENT_EXTENSIONS),
            "ocr": "RapidOCR ONNX CPU with bounded Poppler/FFmpeg rendering",
            "media": "FFprobe streams plus Pillow image metadata and EXIF",
            "archive_formats": ["zip", "7z", "tar", "tgz", "tar.gz", "tar.bz2", "tar.xz"],
            "transcription": {
                "engine": "faster-whisper base CPU int8",
                "model_cached": WHISPER_CACHE.is_dir() and any(WHISPER_CACHE.rglob("model.bin")),
                "maximum_seconds": MAX_TRANSCRIPTION_SECONDS,
            },
            "output_root": str(OUTPUT_ROOT),
            "approved_roots": [str(root) for root in self.policy.roots],
            "boundaries": {
                "authenticated_admin_only": True,
                "public_website": False,
                "roundtable": False,
                "arbitrary_shell": False,
                "automatic_delete_or_commit": False,
            },
        }

    def hash_file(self, path: str) -> dict[str, Any]:
        source = _bounded_file(self.policy, path, max_bytes=20_000_000_000)
        return {"path": str(source), "bytes": source.stat().st_size, "sha256": _sha256_path(source)}

    def search(self, root: str, query: str, mode: str = "content") -> dict[str, Any]:
        target = self.policy.resolve(root)
        if not target.is_dir():
            raise ValueError("search root is not a directory")
        needle = query.strip()
        if not needle or len(needle) > 500:
            raise ValueError("query must contain 1-500 characters")
        normalized_mode = mode.strip().lower()
        if normalized_mode == "content":
            argv = [str(_ripgrep_path()), "--line-number", "--fixed-strings", "--no-messages", "--max-filesize", "20M", "--glob", "!.git/**"]
            for part in sorted(BLOCKED_PARTS):
                argv.extend(["--glob", f"!{part}/**"])
            for name in sorted(BLOCKED_NAMES):
                argv.extend(["--glob", f"!{name}"])
            argv.extend(["--glob", "!.env.*"])
            argv.extend(["--", needle, str(target)])
            completed = subprocess.run(argv, capture_output=True, text=True, timeout=60, env=MINIMAL_COMMAND_ENV)
            if completed.returncode not in {0, 1}:
                raise RuntimeError(completed.stderr.strip()[:1000] or "ripgrep search failed")
            lines = completed.stdout.splitlines()
            return {"root": str(target), "query": needle, "mode": "content", "matches": lines[:MAX_SEARCH_RESULTS], "limited": len(lines) > MAX_SEARCH_RESULTS}
        if normalized_mode not in {"name", "metadata"}:
            raise ValueError("search mode must be content, name, or metadata")
        matches = []
        for directory, names, files in os.walk(target, followlinks=False):
            names[:] = [name for name in names if name not in BLOCKED_PARTS and name != ".git"]
            for name in files:
                if blocked_name(name):
                    continue
                path = Path(directory) / name
                if needle.casefold() not in name.casefold() and normalized_mode == "name":
                    continue
                stat = path.stat()
                record = {"path": str(path), "bytes": stat.st_size, "modified": stat.st_mtime, "extension": path.suffix.casefold()}
                if normalized_mode == "metadata" and needle.casefold() not in json.dumps(record).casefold():
                    continue
                matches.append(record)
                if len(matches) >= MAX_SEARCH_RESULTS:
                    return {"root": str(target), "query": needle, "mode": normalized_mode, "matches": matches, "limited": True}
        return {"root": str(target), "query": needle, "mode": normalized_mode, "matches": matches, "limited": False}

    def duplicates(self, root: str) -> dict[str, Any]:
        target = self.policy.resolve(root)
        if not target.is_dir():
            raise ValueError("duplicate-search root is not a directory")
        by_size: dict[int, list[Path]] = defaultdict(list)
        scanned = 0
        for directory, names, files in os.walk(target, followlinks=False):
            names[:] = [name for name in names if name not in BLOCKED_PARTS and name != ".git"]
            for name in files:
                if blocked_name(name):
                    continue
                path = Path(directory) / name
                try:
                    size = path.stat().st_size
                except OSError:
                    continue
                scanned += 1
                if scanned > MAX_DUPLICATE_FILES:
                    raise ValueError("duplicate search exceeds the 50,000-file scan limit")
                if size:
                    by_size[size].append(path)
        by_digest: dict[str, list[str]] = defaultdict(list)
        for size, paths in by_size.items():
            if len(paths) < 2:
                continue
            for path in paths:
                by_digest[f"{size}:{_sha256_path(path)}"].append(str(path))
        groups = [{"bytes": int(key.split(":", 1)[0]), "sha256": key.split(":", 1)[1], "paths": paths} for key, paths in by_digest.items() if len(paths) > 1]
        return {"root": str(target), "files_scanned": scanned, "duplicate_groups": groups[:200], "groups_limited": len(groups) > 200}


class GitWorkspaceRuntime:
    REPOSITORY = Path("/home/darren/Documents/Little-Oogway")

    def __init__(self, policy: PathPolicy | None = None) -> None:
        self.policy = policy or PathPolicy()

    def _repo(self, value: str) -> Path:
        repo = self.policy.resolve(value)
        if not repo.is_dir() or not (repo / ".git").is_dir():
            raise ValueError("path is not a Git working tree")
        return repo

    @staticmethod
    def _git(repo: Path, arguments: list[str], timeout: int = 60) -> dict[str, Any]:
        completed = subprocess.run(
            [
                "/usr/bin/git", "-c", "color.ui=false", "-c", "core.pager=cat",
                "-c", "core.fsmonitor=false", "-c", "core.hooksPath=/dev/null", *arguments,
            ],
            cwd=repo, capture_output=True, text=True, timeout=timeout, env=MINIMAL_COMMAND_ENV,
        )
        stdout, out_truncated = _trim(completed.stdout, 100_000)
        stderr, err_truncated = _trim(completed.stderr, 20_000)
        return {"exit_code": completed.returncode, "stdout": stdout, "stderr": stderr, "truncated": out_truncated or err_truncated}

    def status(self, repository: str) -> dict[str, Any]:
        repo = self._repo(repository)
        return {"repository": str(repo), **self._git(repo, ["status", "--short", "--branch"])}

    def diff(self, repository: str, staged: bool = False) -> dict[str, Any]:
        repo = self._repo(repository)
        arguments = ["diff", "--no-ext-diff", "--no-textconv", "--no-color"]
        if staged:
            arguments.append("--cached")
        return {"repository": str(repo), "staged": staged, **self._git(repo, arguments)}

    def test(self, repository: str, profile: str = "all") -> dict[str, Any]:
        repo = self.policy.resolve(repository)
        if not repo.is_dir() or repo != self.REPOSITORY:
            raise PermissionError("no bounded test profile is registered for this repository")
        profiles = {
            "core": [([str(repo / "oogway-core/.venv/bin/python"), "-m", "unittest", "discover", "-s", "tests", "-q"], repo / "oogway-core", {"PYTHONPATH": str(repo / "oogway-core/src")})],
            "roundtable": [(["/usr/bin/python3", "-m", "unittest", "discover", "-s", "tests", "-q"], repo / "roundtable", {"PYTHONPATH": str(repo / "roundtable/src")})],
            "all": [
                (["/usr/bin/python3", "-m", "unittest", "discover", "-s", "tokenspace/tests", "-q"], repo, {}),
                ([str(repo / "oogway-core/.venv/bin/python"), "-m", "unittest", "discover", "-s", "tests", "-q"], repo / "oogway-core", {"PYTHONPATH": str(repo / "oogway-core/src")}),
                (["/usr/bin/python3", "-m", "unittest", "discover", "-s", "tests", "-q"], repo / "roundtable", {"PYTHONPATH": str(repo / "roundtable/src")}),
                (["/usr/bin/python3", "-m", "unittest", "discover", "-s", "integrations/public-website", "-q"], repo, {}),
            ],
        }
        selected = profile.strip().lower()
        if selected not in profiles:
            raise ValueError("test profile must be core, roundtable, or all")
        results = []
        for argv, cwd, additions in profiles[selected]:
            environment = dict(MINIMAL_COMMAND_ENV)
            environment.update(additions)
            completed = subprocess.run(argv, cwd=cwd, capture_output=True, text=True, timeout=900, env=environment)
            output, truncated = _trim((completed.stdout + completed.stderr).strip(), 50_000)
            results.append({"argv": argv, "cwd": str(cwd), "exit_code": completed.returncode, "output": output, "truncated": truncated})
            if completed.returncode:
                break
        return {"repository": str(repo), "profile": selected, "passed": bool(results) and all(item["exit_code"] == 0 for item in results), "results": results}

    def propose_commit(self, repository: str) -> dict[str, Any]:
        repo = self._repo(repository)
        status = self._git(repo, ["status", "--short"])
        staged = self._git(repo, ["diff", "--cached", "--stat"])
        unstaged = self._git(repo, ["diff", "--stat"])
        return {
            "repository": str(repo), "status": status["stdout"],
            "staged_summary": staged["stdout"], "unstaged_summary": unstaged["stdout"],
            "commit_executed": False,
            "guidance": "Review the diff, run a registered test profile, stage intentionally, then ask a human or Codex to approve and create the commit.",
        }


class CreativeRuntime:
    """Fixed-profile creative jobs; never accepts command lines or scripts."""

    MEDIA_PROFILES = {
        "web-video": {
            "suffix": ".mp4",
            "arguments": [
                "-map", "0:v:0", "-map", "0:a?", "-vf",
                "scale=w=min(1920\\,iw):h=-2", "-c:v", "libx264", "-preset",
                "medium", "-crf", "22", "-pix_fmt", "yuv420p", "-c:a", "aac",
                "-b:a", "192k", "-movflags", "+faststart",
            ],
        },
        "web-audio": {
            "suffix": ".mp3",
            "arguments": ["-vn", "-c:a", "libmp3lame", "-b:a", "192k"],
        },
        "lossless-audio": {
            "suffix": ".flac",
            "arguments": ["-vn", "-c:a", "flac"],
        },
        "thumbnail": {
            "suffix": ".jpg",
            "arguments": [
                "-frames:v", "1", "-vf", "scale=w=min(1600\\,iw):h=-2",
                "-q:v", "2",
            ],
        },
        "gif-preview": {
            "suffix": ".gif",
            "arguments": [
                "-t", "12", "-vf", "fps=12,scale=w=min(720\\,iw):h=-2:flags=lanczos",
            ],
        },
    }

    def __init__(self, policy: PathPolicy | None = None) -> None:
        self.policy = policy or PathPolicy()

    @staticmethod
    def _provider(name: str) -> str | None:
        candidate = Path("/usr/bin") / name
        return str(candidate) if candidate.is_file() else None

    def capabilities(self) -> dict[str, Any]:
        providers = {
            name: self._provider(name)
            for name in (
                "blender", "obs", "audacity", "kdenlive", "inkscape", "ffmpeg",
                "ffprobe", "pandoc", "tesseract", "sox", "rubberband", "mediainfo",
            )
        }
        return {
            "providers": providers,
            "media_profiles": sorted(self.MEDIA_PROFILES),
            "blender": {
                "headless_render": bool(providers["blender"]),
                "embedded_python_disabled": True,
                "maximum_frames_per_job": 300,
            },
            "output_root": str(OUTPUT_ROOT / "creative"),
            "boundaries": [
                "No arbitrary command lines or application scripts",
                "No OBS streaming control",
                "No Kdenlive or Audacity GUI automation",
                "Outputs are written only below the workbench creative directory",
            ],
        }

    def transcode(self, path: str, profile: str) -> dict[str, Any]:
        source = _bounded_file(self.policy, path, max_bytes=20_000_000_000)
        selected = profile.strip().lower()
        settings = self.MEDIA_PROFILES.get(selected)
        if settings is None:
            raise ValueError(
                "profile must be web-video, web-audio, lossless-audio, thumbnail, or gif-preview"
            )
        if self._provider("ffmpeg") is None:
            raise RuntimeError("FFmpeg is not installed")
        creative_root = OUTPUT_ROOT / "creative"
        creative_root.mkdir(parents=True, exist_ok=True)
        output = _safe_output_name(source, f"-{selected}{settings['suffix']}")
        output = creative_root / output.name
        if output.is_file():
            return {
                "source": str(source), "profile": selected, "output": str(output),
                "reused": True, "sha256": _sha256_path(output),
            }
        temporary = output.with_name(f"{output.stem}.partial{output.suffix}")
        completed = subprocess.run(
            [
                "/usr/bin/ffmpeg", "-nostdin", "-hide_banner", "-loglevel", "warning",
                "-i", str(source), *settings["arguments"], str(temporary),
            ],
            capture_output=True,
            text=True,
            timeout=3_600,
            env=MINIMAL_COMMAND_ENV,
        )
        stderr, truncated = _trim(completed.stderr, 20_000)
        if completed.returncode or not temporary.is_file():
            temporary.unlink(missing_ok=True)
            raise RuntimeError(stderr or "FFmpeg did not create an output")
        os.replace(temporary, output)
        return {
            "source": str(source), "profile": selected, "output": str(output),
            "reused": False, "sha256": _sha256_path(output),
            "log": stderr, "log_truncated": truncated,
        }

    def render_blender(self, path: str, start_frame: int = 1, end_frame: int = 1) -> dict[str, Any]:
        source = _bounded_file(self.policy, path, max_bytes=5_000_000_000)
        if source.suffix.casefold() != ".blend":
            raise ValueError("Blender rendering requires a .blend project")
        if start_frame < 1 or end_frame < start_frame or end_frame - start_frame + 1 > 300:
            raise ValueError("Blender frame range must contain 1-300 frames")
        if self._provider("blender") is None:
            raise RuntimeError("Blender is not installed")
        creative_root = OUTPUT_ROOT / "creative"
        creative_root.mkdir(parents=True, exist_ok=True)
        fingerprint = hashlib.sha256(
            f"{source}:{start_frame}:{end_frame}".encode("utf-8")
        ).hexdigest()[:12]
        output = creative_root / f"{source.stem}-{fingerprint}-blend-render"
        if output.is_dir() and any(output.iterdir()):
            return {
                "source": str(source), "output": str(output), "reused": True,
                "frames": end_frame - start_frame + 1,
            }
        with tempfile.TemporaryDirectory(prefix="oogway-blender-", dir=creative_root) as directory:
            staging = Path(directory)
            completed = subprocess.run(
                [
                    "/usr/bin/blender", "--background", "--factory-startup",
                    "--disable-autoexec", str(source), "--render-output",
                    str(staging / "frame_#####"), "--frame-start", str(start_frame),
                    "--frame-end", str(end_frame), "--render-anim",
                ],
                capture_output=True,
                text=True,
                timeout=7_200,
                env=MINIMAL_COMMAND_ENV,
            )
            log, truncated = _trim(completed.stdout + completed.stderr, 40_000)
            if completed.returncode or not any(staging.iterdir()):
                raise RuntimeError(log or "Blender did not create rendered frames")
            os.replace(staging, output)
        return {
            "source": str(source), "output": str(output), "reused": False,
            "frames": end_frame - start_frame + 1, "log": log,
            "log_truncated": truncated, "embedded_python_disabled": True,
        }
