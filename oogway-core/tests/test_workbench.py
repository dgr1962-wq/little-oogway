from __future__ import annotations

import base64
import json
import unittest
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from oogway_core.operator import PathPolicy
from oogway_core.workbench import (
    ArchiveRuntime,
    DocumentRuntime,
    FileRuntime,
    GitWorkspaceRuntime,
    MediaRuntime,
)


class WorkbenchTests(unittest.TestCase):
    def test_reads_and_converts_markdown_with_provenance(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "notes.md"
            source.write_text("# Notes\n\nBounded tools.\n", encoding="utf-8")
            output = root / "output"
            runtime = DocumentRuntime(PathPolicy((root,)))
            with patch("oogway_core.workbench.OUTPUT_ROOT", output):
                read = runtime.read(str(source))
                converted = runtime.convert_to_markdown(str(source))
            self.assertIn("Bounded tools", read["content"])
            result = Path(converted["output"])
            self.assertTrue(result.is_file())
            self.assertIn(read["sha256"], result.read_text(encoding="utf-8"))

    def test_html_reader_removes_executable_content(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "page.html"
            source.write_text(
                "<html><head><title>Safe title</title><script>attack()</script></head>"
                "<body><h1>Visible</h1><p>Useful text.</p></body></html>",
                encoding="utf-8",
            )
            result = DocumentRuntime(PathPolicy((root,))).read(str(source))
            self.assertIn("Useful text", result["content"])
            self.assertNotIn("attack", result["content"])

    def test_reads_xlsx_spreadsheet_cells(self) -> None:
        from openpyxl import Workbook

        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "budget.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Budget"
            sheet.append(["Item", "Cost"])
            sheet.append(["Quantum cookies", 12.5])
            workbook.save(source)
            result = DocumentRuntime(PathPolicy((root,))).read(str(source))
            self.assertIn("## Sheet: Budget", result["content"])
            self.assertIn("Quantum cookies | 12.5", result["content"])
            self.assertEqual(result["metadata"]["format"], "xlsx")

    def test_reads_ods_spreadsheet_cells(self) -> None:
        from odf.opendocument import OpenDocumentSpreadsheet
        from odf.table import Table, TableCell, TableRow
        from odf.text import P

        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "budget.ods"
            document = OpenDocumentSpreadsheet()
            table = Table(name="Budget")
            for values in (["Item", "Cost"], ["Quantum cookies", "12.5"]):
                row = TableRow()
                for value in values:
                    cell = TableCell(valuetype="string")
                    cell.addElement(P(text=value))
                    row.addElement(cell)
                table.addElement(row)
            document.spreadsheet.addElement(table)
            document.save(str(source))
            result = DocumentRuntime(PathPolicy((root,))).read(str(source))
            self.assertIn("## Sheet: Budget", result["content"])
            self.assertIn("Quantum cookies | 12.5", result["content"])
            self.assertEqual(result["metadata"]["format"], "ods")

    def test_reads_pptx_presentation_text(self) -> None:
        from pptx import Presentation

        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "briefing.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Little Oogway"
            slide.placeholders[1].text = "Bounded local tools"
            presentation.save(source)
            result = DocumentRuntime(PathPolicy((root,))).read(str(source))
            self.assertIn("## Slide 1", result["content"])
            self.assertIn("Bounded local tools", result["content"])
            self.assertEqual(result["metadata"]["format"], "pptx")

    def test_reads_odp_presentation_text(self) -> None:
        from odf.draw import Frame, Page, TextBox
        from odf.opendocument import OpenDocumentPresentation
        from odf.style import MasterPage, PageLayout
        from odf.text import P

        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "briefing.odp"
            document = OpenDocumentPresentation()
            layout = PageLayout(name="Layout")
            document.automaticstyles.addElement(layout)
            document.masterstyles.addElement(MasterPage(name="Default", pagelayoutname="Layout"))
            page = Page(name="slide1", masterpagename="Default")
            frame = Frame(x="1cm", y="1cm", width="20cm", height="10cm")
            box = TextBox()
            box.addElement(P(text="Bounded local tools"))
            frame.addElement(box)
            page.addElement(frame)
            document.presentation.addElement(page)
            document.save(str(source))
            result = DocumentRuntime(PathPolicy((root,))).read(str(source))
            self.assertIn("## Slide 1", result["content"])
            self.assertIn("Bounded local tools", result["content"])
            self.assertEqual(result["metadata"]["format"], "odp")

    def test_legacy_office_conversion_is_isolated_before_parsing(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "old.doc"
            source.write_bytes(b"legacy")
            converted = root / "old.docx"
            temporary = Mock()
            runtime = DocumentRuntime(PathPolicy((root,)))
            with patch.object(runtime, "_convert_legacy_office", return_value=(converted, temporary)), patch.object(
                runtime, "_extract", side_effect=[("converted text", {"format": "docx"})]
            ) as extract:
                text, metadata = DocumentRuntime._extract(runtime, source, ocr_if_needed=True, max_pages=30)
            self.assertEqual(text, "converted text")
            self.assertEqual(metadata["converted_via"], "LibreOffice")
            self.assertEqual(metadata["intermediate_format"], "docx")
            temporary.cleanup.assert_called_once()
            extract.assert_called_once()

    def test_ocr_image_is_bounded_and_reports_confidence(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "page.png"
            source.write_bytes(
                base64.b64decode(
                    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
                )
            )
            runtime = DocumentRuntime(PathPolicy((root,)))
            with patch.object(runtime, "_ocr_image", return_value=("hello page", [0.9, 0.8])):
                result = runtime.ocr(str(source), max_pages=1)
            self.assertEqual(result["content"], "## Page 1\n\nhello page")
            self.assertEqual(result["mean_confidence"], 0.85)

    def test_archive_rejects_path_traversal(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "bad.zip"
            with zipfile.ZipFile(source, "w") as archive:
                archive.writestr("../escape.txt", "no")
            runtime = ArchiveRuntime(PathPolicy((root,)))
            with self.assertRaisesRegex(ValueError, "unsafe archive path"):
                runtime.inspect(str(source))

    def test_archive_extracts_only_into_bounded_output(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "safe.zip"
            with zipfile.ZipFile(source, "w") as archive:
                archive.writestr("folder/readme.txt", "safe")
            runtime = ArchiveRuntime(PathPolicy((root,)))
            output = root / "output"
            with patch("oogway_core.workbench.OUTPUT_ROOT", output):
                result = runtime.extract(str(source))
            self.assertEqual(Path(result["output"]).joinpath("folder/readme.txt").read_text(), "safe")

    def test_duplicate_detection_hashes_only_equal_sizes(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "one.bin").write_bytes(b"same")
            (root / "two.bin").write_bytes(b"same")
            (root / "other.bin").write_bytes(b"different")
            result = FileRuntime(PathPolicy((root,))).duplicates(str(root))
            self.assertEqual(len(result["duplicate_groups"]), 1)
            self.assertEqual(len(result["duplicate_groups"][0]["paths"]), 2)

    def test_capability_inventory_states_security_boundaries(self) -> None:
        with TemporaryDirectory() as directory:
            result = FileRuntime(PathPolicy((Path(directory),))).capabilities()
            self.assertTrue(result["boundaries"]["authenticated_admin_only"])
            self.assertFalse(result["boundaries"]["public_website"])
            self.assertFalse(result["boundaries"]["arbitrary_shell"])

    def test_media_inspection_uses_structured_ffprobe(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "clip.mp3"
            source.write_bytes(b"audio")
            completed = Mock(returncode=0, stdout=json.dumps({"format": {"duration": "1.0"}}))
            with patch("oogway_core.workbench.subprocess.run", return_value=completed):
                result = MediaRuntime(PathPolicy((root,))).inspect(str(source))
            self.assertEqual(result["ffprobe"]["format"]["duration"], "1.0")

    def test_git_tests_fail_closed_for_unregistered_repository(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / ".git").mkdir()
            runtime = GitWorkspaceRuntime(PathPolicy((root,)))
            with self.assertRaisesRegex(PermissionError, "no bounded test profile"):
                runtime.test(str(root), "all")


if __name__ == "__main__":
    unittest.main()
